from docx import Document
from server.extraction import local_pages

def test_docx_extract_preserves_table_text(tmp_path):
    source=tmp_path/'synthetic.docx'
    doc=Document();doc.add_paragraph('合成测试表单');doc.add_table(rows=1,cols=1).cell(0,0).text='申请人签字';doc.save(source)
    pages=local_pages(source)
    assert pages and '申请人签字' in pages[0]['text']
    assert pages[0]['page_label']=='Word 正文（未分页）'
